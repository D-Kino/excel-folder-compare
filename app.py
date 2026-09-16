"""Offline Excel folder comparison. Does not automate or launch Excel."""
from __future__ import annotations
import argparse
import datetime as dt
import math
import os
from pathlib import Path
import queue
import threading
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

SUPPORTED = {'.xlsx', '.xlsm'}
KNOWN = SUPPORTED | {'.xls', '.xlsb'}
BLANK = ('空白', '', '', 'なし', '空白', '')

def value_info(value):
    if value is None: return ('空白', '')
    if isinstance(value, bool): return ('真偽値', 'TRUE' if value else 'FALSE')
    if isinstance(value, (dt.datetime, dt.date, dt.time)): return ('日時', value.isoformat())
    if isinstance(value, dt.timedelta): return ('時間', str(value))
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer(): value = int(value)
        return ('数値', str(value))
    return ('文字列', str(value))

def strike_info(cell):
    value = cell.value
    if not isinstance(value, CellRichText): return '全体' if cell.font.strike else 'なし'
    flags = []
    for part in value:
        if isinstance(part, TextBlock):
            flags.extend([bool(part.font.strike)] * len(part.text))
        else:
            flags.extend([bool(cell.font.strike)] * len(str(part)))
    if not flags: return '全体' if cell.font.strike else 'なし'
    if all(flags): return '全体'
    if not any(flags): return 'なし'
    ranges, start = [], None
    for i, on in enumerate(flags + [False], 1):
        if on and start is None: start = i
        if not on and start is not None:
            ranges.append(str(start) if start == i-1 else f'{start}-{i-1}')
            start = None
    return '一部（文字位置: ' + ', '.join(ranges) + '）'

def snapshot(path):
    book = cached = None
    try:
        book = load_workbook(path, rich_text=True, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        sheets, notes = {}, []
        for sheet in book.worksheets:
            cells = {}
            # Sparse enumeration includes styled empty cells without iterating a huge blank rectangle.
            for (row, col), cell in sorted(sheet._cells.items()):
                value = cell.value
                strike = strike_info(cell)
                formula = ''
                if cell.data_type == 'f':
                    if isinstance(value, ArrayFormula):
                        formula = f'{value.text} [配列範囲:{value.ref}]'
                    elif isinstance(value, DataTableFormula):
                        raise ValueError(f'{sheet.title}!{cell.coordinate}: What-Ifデータテーブルの数式は比較できません。')
                    else: formula = str(value)
                    kind, text = ('数式', '')
                    stored = cached[sheet.title].cell(row, col)
                    ck, ct = value_info(stored.value)
                    if stored.data_type == 'e': ck = 'Excelエラー'
                    if stored.value is None: ck, ct = '保存値なし', ''
                else:
                    kind, text = value_info(value)
                    if cell.data_type == 'e': kind = 'Excelエラー'
                    ck, ct = '空白', ''
                if value is not None or strike != 'なし':
                    cells[(row, col)] = (kind, text, formula, strike, ck, ct)
            for rules in sheet.conditional_formatting._cf_rules.values():
                if any(rule.dxf and rule.dxf.font and rule.dxf.font.strike is not None for rule in rules):
                    notes.append(f'{sheet.title}: 条件付き書式による取消線は対象外（直接設定された取消線を比較）')
                    break
            if any(dim.font.strike for dim in list(sheet.row_dimensions.values()) + list(sheet.column_dimensions.values())):
                notes.append(f'{sheet.title}: 行・列全体の取消線設定あり。セルに保存された書式のみ比較')
            sheets[sheet.title] = cells
        return sheets, notes
    finally:
        if book: book.close()
        if cached: cached.close()

def files_in(folder, output):
    result = {}
    for path in sorted(folder.iterdir()):
        if path.is_file() and path.suffix.lower() in KNOWN and not path.name.startswith('~$') and path.resolve() != output:
            key = path.name.casefold()
            if key in result: raise ValueError(f'同名と判断されるファイルが重複しています: {path.name}')
            result[key] = path
    return result

def text_cell(sheet, row, col, value):
    cell = sheet.cell(row, col)
    if isinstance(value, int): cell.value = value
    else:
        text = str(value)
        if len(text) > 32767: raise ValueError('結果の1セルがExcelの文字数上限を超えました。')
        cell.value = text
        cell.data_type = 's'  # Never execute a source formula, including =, +, @ or - text.
    return cell

def add_table(book, title, headers, rows, widths):
    sheet = book.create_sheet(title)
    if len(rows) >= 1048576: raise ValueError('差分がExcelの行数上限を超えました。入力フォルダを分けてください。')
    for ri, record in enumerate([headers] + rows, 1):
        for ci, value in enumerate(record, 1):
            cell = text_cell(sheet, ri, ci, value)
            cell.font = Font(name='Yu Gothic', size=10, color='FFFFFF' if ri == 1 else '203047', bold=ri == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if ri == 1: cell.fill = PatternFill('solid', fgColor='234B68')
            elif ri % 2 == 0: cell.fill = PatternFill('solid', fgColor='F0F5F8')
        # Estimate readable row height; Excel supports at most 409 points.
        lines = max((sum(max(1, math.ceil(len(line) / max(8, widths[min(ci, len(widths)-1)] / 2))) for line in str(v).split('\n')) for ci, v in enumerate(record)), default=1)
        sheet.row_dimensions[ri].height = min(409, max(30, 16 * lines + 8))
    for i, width in enumerate(widths, 1): sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    return sheet

def compare(folder_a, folder_b, output, progress=lambda message: None):
    a, b, out = Path(folder_a).resolve(), Path(folder_b).resolve(), Path(output).resolve()
    if not a.is_dir() or not b.is_dir(): raise ValueError('A・Bには存在するフォルダを指定してください。')
    if a == b: raise ValueError('A・Bには別々のフォルダを指定してください。')
    if out.suffix.lower() != '.xlsx': raise ValueError('出力名は.xlsxにしてください。')
    if out.exists(): raise ValueError('出力先のファイルは既に存在します。新しい名前を指定してください。')
    if not out.parent.is_dir(): raise ValueError('出力先のフォルダが存在しません。')
    fa, fb = files_in(a, out), files_in(b, out)
    if not fa and not fb: raise ValueError('選択したフォルダ直下にExcelファイルがありません。')
    summaries, details = [], []
    for index, key in enumerate(sorted(fa.keys() | fb.keys()), 1):
        pa, pb = fa.get(key), fb.get(key)
        name = (pa or pb).name
        progress(f'{index}/{len(fa.keys() | fb.keys())}  {name}')
        if (pa or pb).suffix.lower() not in SUPPORTED:
            summaries.append([name, '対象外', '', '.xlsxまたは.xlsmに変換してください。']); continue
        if pa is None or pb is None:
            summaries.append([name, 'Bのみ' if pa is None else 'Aのみ', '', '同名ファイルがありません。']); continue
        start = len(details)
        try:
            sa, na = snapshot(pa); sb, nb = snapshot(pb)
            for sn in sorted(sa.keys() | sb.keys()):
                if sn not in sa or sn not in sb:
                    details.append([name, sn, '', 'Bのみのシート' if sn not in sa else 'Aのみのシート'] + [''] * 10)
                    continue
                for pos in sorted(sa[sn].keys() | sb[sn].keys()):
                    va, vb = sa[sn].get(pos, BLANK), sb[sn].get(pos, BLANK)
                    changes = []
                    if va[:2] != vb[:2]: changes.append('値・型')
                    if va[2] != vb[2]: changes.append('数式')
                    if va[3] != vb[3]: changes.append('取消線')
                    if va[4:] != vb[4:]: changes.append('保存済み計算結果')
                    if changes:
                        from openpyxl.utils import get_column_letter
                        addr = f'{get_column_letter(pos[1])}{pos[0]}'
                        details.append([name, sn, addr, '・'.join(changes), va[1], vb[1], va[2], vb[2], va[3], vb[3], va[0], vb[0], va[5] if va[4] != '保存値なし' else '（保存値なし）', vb[5] if vb[4] != '保存値なし' else '（保存値なし）'])
            count = len(details) - start
            warnings = list(dict.fromkeys(na + nb))
            state = '差分あり' if count else '差分なし'
            if warnings: state += '（注意あり）'
            summaries.append([name, state, count, '\n'.join(warnings)])
        except Exception as exc:
            del details[start:]
            summaries.append([name, 'エラー', '', f'{type(exc).__name__}: {exc}'])
    progress('比較結果を保存しています…')
    book = Workbook(); book.remove(book.active)
    add_table(book, '比較一覧', ['ファイル名', '結果', '差分行数', '備考'], summaries, [35, 25, 15, 90])
    detail_sheet = add_table(book, '差分詳細', ['ファイル名', 'シート名', 'セル', '差分の種類', 'Aの値', 'Bの値', 'Aの数式', 'Bの数式', 'Aの取消線', 'Bの取消線', 'Aの型', 'Bの型', 'Aの保存済み計算結果', 'Bの保存済み計算結果'], details, [28, 20, 12, 26, 36, 36, 36, 36, 30, 30, 15, 15, 30, 30])
    for row in detail_sheet.iter_rows(min_row=2, max_row=len(details)+1, min_col=9, max_col=10):
        for cell in row: cell.fill = PatternFill('solid', fgColor='FFF0CC')
    conditions = [
        ['Aフォルダ', str(a)], ['Bフォルダ', str(b)], ['実行日時', dt.datetime.now().isoformat(timespec='seconds')],
        ['比較方法', '直下の同名ファイル（拡張子込み・大文字小文字を区別しない）→同名シート→同じセル位置。サブフォルダは含みません。'],
        ['比較対象', '値・型、数式、保存済み計算結果、直接設定された取消線（文字の一部を含む）。取消線の文字位置は先頭を1とします。'],
        ['対象外', '色・罫線・その他書式、条件付き書式の表示結果、行列全体の書式継承、画像・グラフ・コメント・マクロ。'],
        ['数式', '再計算はしません。保存済みの計算結果を比較します。結果が未保存の場合は「保存値なし」と表示します。'],
        ['差分行数', '変更セル1件を1行、片方だけのシート1件を1行として集計。行挿入や並べ替えを追跡しません。'],
        ['型・日時', '数値と文字列、空白とゼロを区別します。日付はISO形式で表示。数値の許容誤差は設けません。'],
        ['エラー', '暗号化・破損ファイル、What-Ifデータテーブルなど比較できないファイルは比較一覧に記録します。'],
    ]
    add_table(book, '比較条件', ['項目', '内容'], conditions, [25, 110])
    # Exclusive creation protects inputs and existing reports even if a file appears during processing.
    try:
        with out.open('xb') as handle: book.save(handle)
    except FileExistsError: raise ValueError('出力先のファイルが作成されました。別名で再実行してください。')
    except Exception:
        if out.exists(): out.unlink()
        raise
    finally: book.close()
    return {'files': len(summaries), 'differences': len(details), 'errors': sum(r[1] == 'エラー' for r in summaries), 'unsupported': sum(r[1] == '対象外' for r in summaries), 'output': str(out)}

def gui():
    root = tk.Tk(); root.title('Excelフォルダ比較'); root.geometry('760x355'); root.minsize(640, 355)
    root.option_add('*Font', ('Yu Gothic UI', 10))
    frame = ttk.Frame(root, padding=20); frame.pack(fill='both', expand=True); frame.columnconfigure(1, weight=1)
    ttk.Label(frame, text='Excelファイルの差分を確認', font=('Yu Gothic UI', 16, 'bold')).grid(row=0, column=0, columnspan=3, sticky='w', pady=(0, 8))
    ttk.Label(frame, text='値・数式・取消線を比較します。対応形式: .xlsx / .xlsm').grid(row=1, column=0, columnspan=3, sticky='w', pady=(0, 14))
    variables = [tk.StringVar() for _ in range(3)]
    controls = []
    def browse(i):
        value = filedialog.askdirectory(parent=root, title='Aフォルダ' if i == 0 else 'Bフォルダ') if i < 2 else filedialog.asksaveasfilename(parent=root, title='結果の保存先（新しい名前）', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')], initialfile='Excel差分_' + dt.datetime.now().strftime('%Y%m%d_%H%M%S') + '.xlsx')
        if value: variables[i].set(value)
    for i, label in enumerate(['Aフォルダ', 'Bフォルダ', '結果の保存先']):
        ttk.Label(frame, text=label).grid(row=i+2, column=0, sticky='w', padx=(0, 12), pady=6)
        entry = ttk.Entry(frame, textvariable=variables[i]); entry.grid(row=i+2, column=1, sticky='ew', pady=6)
        button = ttk.Button(frame, text='選択…', command=lambda i=i: browse(i)); button.grid(row=i+2, column=2, padx=(8, 0))
        controls.extend([entry, button])
    status = tk.StringVar(value='フォルダと保存先を指定して「比較する」を押してください。')
    ttk.Label(frame, textvariable=status, wraplength=660).grid(row=6, column=0, columnspan=3, sticky='w', pady=12)
    messages = queue.Queue(); running = [False]
    def worker(args):
        try: messages.put(('done', compare(*args, progress=lambda x: messages.put(('progress', x)))))
        except Exception as exc: messages.put(('error', str(exc)))
    def start():
        args = [v.get().strip() for v in variables]
        if not all(args): messagebox.showerror('指定が必要です', 'A・Bフォルダと保存先をすべて指定してください。', parent=root); return
        running[0] = True
        for widget in controls: widget.configure(state='disabled')
        status.set('比較を開始しています…')
        threading.Thread(target=worker, args=(args,), daemon=True).start()
    button = ttk.Button(frame, text='比較する', command=start); button.grid(row=5, column=2, pady=(12, 0)); controls.append(button)
    def poll():
        try:
            while True:
                kind, data = messages.get_nowait()
                if kind == 'progress': status.set(data)
                else:
                    running[0] = False
                    for widget in controls: widget.configure(state='normal')
                    if kind == 'error': status.set('処理を完了できませんでした。'); messagebox.showerror('Excel比較', data, parent=root)
                    else:
                        text = f"保存完了: {data['files']}ファイル / 差分{data['differences']}行 / エラー{data['errors']}件 / 対象外{data['unsupported']}件"
                        status.set(text); messagebox.showinfo('Excel比較', text + '\n\n' + data['output'], parent=root)
        except queue.Empty: pass
        root.after(150, poll)
    def close():
        if running[0]: messagebox.showinfo('比較中', '保存が完了するまでお待ちください。', parent=root)
        else: root.destroy()
    root.protocol('WM_DELETE_WINDOW', close); poll(); root.mainloop()

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--a'); parser.add_argument('--b'); parser.add_argument('--out')
    args = parser.parse_args()
    if args.a and args.b and args.out: print(compare(args.a, args.b, args.out, print))
    else: gui()
