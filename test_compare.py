from pathlib import Path
import importlib.util
import uuid
import hashlib
import zipfile
import xml.etree.ElementTree as ET
import sys
import tkinter as tk
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.formula import ArrayFormula

root = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('compare_app', root/'app.py')
app = importlib.util.module_from_spec(spec); spec.loader.exec_module(app)
(root/'work').mkdir(exist_ok=True)
testroot = root/'work'/('direct_' + uuid.uuid4().hex[:10]); testroot.mkdir()
a, b = testroot/'A', testroot/'B'; a.mkdir(); b.mkdir()

def fixture(side):
    w = Workbook(); s = w.active; s.title = 'Data'
    s['A1'] = 'same'
    s['A2'] = 'strike'; s['A2'].font = Font(strike=side=='B')
    s['A3'] = CellRichText('a', TextBlock(InlineFont(strike=True), 'bc'), 'd', TextBlock(InlineFont(strike=side=='B'), 'e'), 'f')
    s['A4'] = '=1+1' if side=='A' else '=2+0'
    s['A5'] = 'ABC' if side=='A' else 'abc'
    s['A6'].font = Font(strike=side=='B')
    s['A7'] = '=literal'; s['A7'].data_type = 's'
    s['A8'] = 'same'; s['A8'].font = Font(color='000000' if side=='A' else 'FF0000')
    if side=='A': s['A9'] = 0
    s['A10'] = '1' if side=='A' else 1
    s['A11'] = True if side=='A' else 1
    s['A12'] = 'has space ' if side=='A' else 'has space'
    s['A13'] = 'a\nb' if side=='A' else 'ab'
    s['A14'] = CellRichText(TextBlock(InlineFont(strike=True), 'ab'), TextBlock(InlineFont(strike=True), 'cd')) if side=='A' else CellRichText(TextBlock(InlineFont(strike=True), 'abcd'))
    s['A15'] = '=9+1'
    s['A16'] = '日本語比較'; s['A16'].font = Font(strike=True)
    s['Z1000000'].fill = PatternFill('solid', fgColor='FFFF00')
    s['B1'] = '=A1'; s['B2'] = '=A2'
    w.create_sheet('A_only' if side=='A' else 'B_only')
    w.save((a if side=='A' else b)/'match.xlsx')
    w.close()

for side in ('A','B'): fixture(side)
ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def xml_patch(path, side):
    with zipfile.ZipFile(path) as z: parts = {i.filename:z.read(i) for i in z.infolist()}
    xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    for c in xml.findall('m:sheetData/m:row/m:c',ns):
        if c.attrib['r']=='A15': c.find('m:v',ns).text = '10' if side=='A' else '11'
        if side=='A' and c.attrib['r'] in ('B1','B2'):
            f = c.find('m:f',ns); f.set('t','shared'); f.set('si','0')
            if c.attrib['r']=='B1': f.set('ref','B1:B2')
            else: f.text = None
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items(): z.writestr(name,data)
for side in ('A','B'): xml_patch((a if side=='A' else b)/'match.xlsx',side)
for folder in (a,b):
    (folder/'broken.xlsx').write_text('broken')
    (folder/'legacy.xls').write_text('old')
    (folder/'~$ignored.xlsx').write_text('ignore')
    w=Workbook(); w.active['A1']='identical'; w.save(folder/'same.xlsx'); w.close()
    w=Workbook(); w.active['A1']='conditional'; w.active.conditional_formatting.add('A1',CellIsRule(operator='equal', formula=['1'],font=Font(strike=True))); w.save(folder/'conditional.xlsx'); w.close()
(a/'only_a.xlsx').write_bytes((a/'same.xlsx').read_bytes())
(b/'only_b.xlsx').write_bytes((b/'same.xlsx').read_bytes())
before = {str(p):hashlib.sha256(p.read_bytes()).hexdigest() for folder in (a,b) for p in folder.iterdir()}
result = app.compare(a,b,testroot/'result.xlsx',print)
assert result['errors']==1 and result['unsupported']==1, result
w=load_workbook(testroot/'result.xlsx',data_only=False)
details=list(w['差分詳細'].values)
rows = {r[2]:r for r in details[1:] if r[2]}
assert set(rows)=={'A2','A3','A4','A5','A6','A9','A10','A11','A12','A13','A15'}, rows.keys()
assert rows['A3'][8]=='一部（文字位置: 2-3）' and rows['A3'][9]=='一部（文字位置: 2-3, 5）', rows['A3']
assert rows['A4'][6]=='=1+1' and rows['A4'][7]=='=2+0'
assert rows['A15'][12:]==('10','11'), rows['A15']
assert rows['A2'][3]=='取消線' and rows['A6'][3]=='取消線'
assert result['differences']==13, result
summary={r[0]:r for r in list(w['比較一覧'].values)[1:]}
assert summary['same.xlsx'][1]=='差分なし'
assert summary['conditional.xlsx'][1]=='差分なし（注意あり）'
assert summary['only_a.xlsx'][1]=='Aのみ' and summary['only_b.xlsx'][1]=='Bのみ'
for sheet in w:
    assert sheet.freeze_panes=='A2' and sheet.auto_filter.ref
    assert not any(c.data_type=='f' for row in sheet for c in row), 'Output formula executed'
w.close()
assert before=={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for folder in (a,b) for p in folder.iterdir()}
try: app.compare(a,b,testroot/'result.xlsx')
except ValueError: pass
else: raise AssertionError('Overwrite permitted')
# Exercise the actual GUI startup with the bundled Tcl/Tk runtime and stop automatically.
original_mainloop = tk.Tk.mainloop
def short_loop(self, *args, **kwargs):
    self.after(200, self.destroy)
    original_mainloop(self, *args, **kwargs)
tk.Tk.mainloop=short_loop
app.gui()
print('PASS: whole/partial/empty strikes, values, types, formulas, shared formulas, cached values, sparse sheets, conditional warning, missing files/sheets, broken/unsupported files, no overwrite, source hashes, report reopen, GUI startup')
print('REPORT:',testroot/'result.xlsx')
